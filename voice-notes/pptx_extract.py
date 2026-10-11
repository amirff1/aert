#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
استخراج محتوای پاورپوینت (.pptx)
==================================
* متن هر اسلاید (عنوان + بدنه + جدول‌ها)
* نُت‌های گوینده (speaker notes)
* فهرست مدیای جاسازی‌شده (صدا/ویدیو)

خروجی:  slides.md  و  slides.json  در پوشهٔ مقصد.

استفاده:
    python3 pptx_extract.py FILE.pptx [--out DIR]
"""
from __future__ import annotations

import argparse
import json
import zipfile
from pathlib import Path


def _shape_text(shape) -> list[str]:
    out: list[str] = []
    if shape.shape_type == 6 or getattr(shape, "shapes", None):  # group
        for sub in shape.shapes:
            out.extend(_shape_text(sub))
        return out
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            txt = "".join(r.text for r in p.runs).strip()
            if txt:
                out.append(("• " * 1 if not p.level else "  " * p.level + "◦ ") + txt)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                out.append("| " + " | ".join(cells) + " |")
    return out


def extract_pptx(path: Path, out_dir: Path) -> dict:
    from pptx import Presentation

    out_dir.mkdir(parents=True, exist_ok=True)
    prs = Presentation(str(path))
    slides: list[dict] = []

    for i, slide in enumerate(prs.slides, 1):
        title = ""
        try:
            if slide.shapes.title is not None:
                title = slide.shapes.title.text.strip()
        except Exception:
            pass
        body: list[str] = []
        for shape in slide.shapes:
            if shape == getattr(slide.shapes, "title", None):
                continue
            if shape.has_text_frame and shape.text.strip() == title:
                continue
            body.extend(_shape_text(shape))

        notes = ""
        if slide.has_notes_slide:
            notes = "\n".join(
                p.text.strip() for p in slide.notes_slide.notes_text_frame.paragraphs if p.text.strip()
            )

        slides.append({"index": i, "title": title, "body": body, "notes": notes})

    # مدیای جاسازی‌شده
    media: list[str] = []
    try:
        with zipfile.ZipFile(path) as z:
            for n in z.namelist():
                if "/media/" in n and Path(n).suffix.lower() in {
                    ".mp3", ".wav", ".m4a", ".aac", ".ogg", ".opus", ".flac", ".wma",
                    ".mp4", ".mov", ".mkv", ".webm", ".avi", ".wmv", ".m4v",
                }:
                    media.append(n)
    except zipfile.BadZipFile:
        pass

    info = {"file": path.name, "n_slides": len(slides), "media": sorted(media), "slides": slides}
    (out_dir / "slides.json").write_text(json.dumps(info, ensure_ascii=False, indent=2), encoding="utf-8")

    md = [f"# {path.stem}", ""]
    for s in slides:
        md.append(f"## اسلاید {s['index']}" + (f" — {s['title']}" if s["title"] else ""))
        md.append("")
        for b in s["body"]:
            md.append(b)
        if s["notes"]:
            md.append("")
            md.append(f"> 🗒️ نُت گوینده: {s['notes']}")
        md.append("")
    if media:
        md.append("## مدیای جاسازی‌شده")
        for m in sorted(media):
            md.append(f"- `{m}`")
    (out_dir / "slides.md").write_text("\n".join(md), encoding="utf-8")
    return info


def main() -> None:
    ap = argparse.ArgumentParser(description="استخراج متن/نُت/مدیای پاورپوینت")
    ap.add_argument("pptx")
    ap.add_argument("--out", default="slides")
    args = ap.parse_args()
    p = Path(args.pptx).expanduser()
    info = extract_pptx(p, Path(args.out).expanduser())
    print(f"✓ {info['n_slides']} اسلاید، {len(info['media'])} مدیا → {args.out}/slides.md")
    for s in info["slides"]:
        head = (s["title"] or (s["body"][0][:60] if s["body"] else ""))[:70]
        print(f"  اسلاید {s['index']:2d}: {head}")


if __name__ == "__main__":
    main()
